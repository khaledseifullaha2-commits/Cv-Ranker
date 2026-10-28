#!/usr/bin/env python3
"""
tests/smoke_app.py -- the security + pipeline check for app.py
==============================================================

This is the evidence behind the claims in SECURITY.md §2.  Run it after any
change to app.py, evaluator.py or a dependency bump:

    python tests/smoke_app.py

It is deliberately self-sufficient so it can be published with the repository:

* job description -> `1_Active_JD/current_jd.txt`, or the synthetic
  `web/public/demo/jd.txt` when the real posting is absent;
* candidate CVs   -> real files from `2_Active_CVs/`, or two synthetic PDFs
  generated in-process when that folder is absent (CI, a fresh clone);
* blacklist       -> `blacklist.csv` when present, otherwise an in-memory
  entry, so the phone/name behaviour is still exercised without shipping
  anyone's personal data.

No real candidate data is read unconditionally, and nothing is written outside
a temporary directory.
"""
from __future__ import annotations

import copy
import dataclasses
import io
import os
import sys
import types
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
os.chdir(ROOT)
sys.path.insert(0, str(ROOT))          # same sys.path app.py gets when run directly

from streamlit.testing.v1 import AppTest  # noqa: E402  (imported early: it is slow)

# ── load app.py's helper section (everything above the sidebar wiring) ────
_src = (ROOT / "app.py").read_text(encoding="utf-8")
_cut = _src.index("# SECTION 4 -- THE SIDEBAR")
_app = types.ModuleType("app_helpers")          # registered so dataclasses resolve
_app.__file__ = str(ROOT / "app.py")
sys.modules["app_helpers"] = _app
exec(compile(_src[:_cut], "app.py", "exec"), _app.__dict__)

safe_filename = _app.safe_filename
validate_uploads = _app.validate_uploads
read_upload = _app.read_upload
read_pdf_bytes = _app.read_pdf_bytes
read_docx_bytes = _app.read_docx_bytes
keyword_filters = _app.keyword_filters
bonus_points = _app.bonus_points
recruiter_remark = _app.recruiter_remark
exclusion_remark = _app.exclusion_remark
error_remark = _app.error_remark
education_verdict = _app.education_verdict
export_frame = _app.export_frame
build_profile = _app.build_profile
split_keywords = _app.split_keywords
CVReadError = _app.CVReadError
MAX_FILE_MB = _app.MAX_FILE_MB
ev = _app.ev

fails: list[str] = []


def check(name: str, cond: bool, detail: str = "") -> None:
    print(("  PASS  " if cond else "  FAIL  ") + name + ("  [%s]" % detail if detail else ""))
    if not cond:
        fails.append(name)


# ── fixtures ──────────────────────────────────────────────────────────────
class Up:
    """Stand-in for Streamlit's UploadedFile."""

    def __init__(self, name: str, data: bytes, kind: str | None = None):
        self.name, self.data = name, data
        self.kind = kind or ("docx" if name.lower().endswith(".docx") else "pdf")

    def getvalue(self) -> bytes:
        return self.data

    @property
    def size(self) -> int:
        return len(self.data)


def _pdf_escape(text: str) -> str:
    return text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")


def synthetic_cv_pdf(lines: list[str]) -> bytes:
    """A minimal, valid one-page PDF (Helvetica) built without any dependency."""
    stream = "BT /F1 11 Tf 14 TL 50 780 Td\n"
    stream += "".join("(%s) Tj T*\n" % _pdf_escape(line) for line in lines)
    stream += "ET"
    body = stream.encode("latin-1", "replace")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length %d >>\nstream\n" % len(body) + body + b"\nendstream",
    ]

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + obj + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n" % (len(objects) + 1)
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += (b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
            % (len(objects) + 1, xref))
    return bytes(out)


STRONG_CV = synthetic_cv_pdf([
    "MD. RAHIM UDDIN",
    "Phone: 01711-000111",
    "Site Management Field Assistant",
    "",
    "OBJECTIVE",
    "Site management professional with 4 years of experience in camp",
    "coordination, community mobilization and site management activities",
    "in Cox's Bazar humanitarian response.",
    "",
    "EXPERIENCE",
    "1. Site Management Officer, NGO, Ukhiya, Cox's Bazar    (01/2023 - present)",
    "2. Community Mobilizer, NGO, Cox's Bazar                (02/2021 - 12/2022)",
    "3. Field Assistant, Consulting Firm, Dhaka               (01/2020 - 01/2021)",
    "",
    "EDUCATION",
    "Bachelor of Social Science (BSS), University of Chittagong, 2019",
    "",
    "SKILLS",
    "community mobilization, capacity building, data collection,",
    "report writing, teamwork, coordination, leadership",
    "Languages: Bangla (fluent), English (fluent)",
])

WEAK_CV = synthetic_cv_pdf([
    "KARIM SHEIKH",
    "Phone: 01822-000222",
    "ACCOUNTANT",
    "",
    "EXPERIENCE",
    "1. Accounts Assistant, Garment Factory, Gazipur   (03/2019 - present)",
    "",
    "EDUCATION",
    "Secondary school diploma (SSC), 2015",
    "",
    "SKILLS",
    "Tally, MS Excel, payroll, invoicing, book keeping",
    "Languages: Bangla",
])


def cv_fixtures(n: int = 6) -> list[Up]:
    """Real CVs when the folder exists, synthetic ones otherwise."""
    folder = ROOT / "2_Active_CVs"
    real = sorted(folder.glob("*.pdf")) if folder.is_dir() else []
    if real:
        return [Up(p.name, p.read_bytes()) for p in real[:n]]
    return [Up("md-rahim-uddin.pdf", STRONG_CV), Up("karim-sheikh.pdf", WEAK_CV)]


def jd_text() -> str:
    for candidate in (ROOT / "1_Active_JD" / "current_jd.txt",
                      ROOT / "web" / "public" / "demo" / "jd.txt"):
        if candidate.is_file():
            return candidate.read_text(encoding="utf-8", errors="replace")
    raise SystemExit("no job description found")


def blacklist_entries():
    """The real blacklist when present, otherwise an equivalent in-memory one."""
    entries = ev.load_blacklist()
    if entries:
        return entries, True
    # Build the fallback exactly the way load_blacklist() would: names through
    # _name_variants(), phones through _norm_phone() (last-10-digit
    # fingerprint).  Hand-writing the raw values silently breaks matching,
    # which is precisely the bug this fixture once had.
    return [ev.BlacklistEntry(label="Synthetic Test Entry",
                              names=ev._name_variants("Synthetic Test Entry"),
                              phones=[ev._norm_phone("01700000099")])], False


# ══════════════════════════════════════════════════════════════════════════
print("\n=== 1. path-traversal / filename sanitisation ===")
cases = {
    "../../etc/passwd.pdf": "passwd.pdf",
    "..\\..\\Windows\\system32\\evil.pdf": "evil.pdf",
    "C:\\Users\\x\\CVs\\Jane Doe.pdf": "Jane Doe.pdf",
    "..%2F..%2Fetc%2Fshadow.pdf": "shadow.pdf",
    "a/b/c/../../../final.pdf": "final.pdf",
    "résumé<script>alert(1)</script>.pdf": None,      # must merely be safe
    "normal candidate.pdf": "normal candidate.pdf",
    "no_extension": "no extension.pdf",
    "evil.exe": "evil.pdf",                            # bad ext replaced, not kept
}
for raw, want in cases.items():
    got = safe_filename(raw)
    ok = ("/" not in got) and ("\\" not in got) and (":" not in got) \
        and (".." not in got) and ("<" not in got) and ("\x00" not in got) \
        and got.lower().endswith(".pdf")
    if want is not None:
        ok = ok and got == want
    check("safe_filename(%r) -> %r" % (raw, got), ok, want or "safe")

print("\n=== 2. upload validation (magic bytes, size, type) ===")
real_pdf = cv_fixtures(1)[0].data

_, rej = validate_uploads([])
check("no files -> no rejects", rej == [])

_, rej = validate_uploads([Up("empty.pdf", b"")])
check("empty file rejected", len(rej) == 1 and "empty" in rej[0])

_, rej = validate_uploads([Up("fake.pdf", b"<html>this is not a pdf</html>")])
check("non-PDF bytes with .pdf name rejected", len(rej) == 1 and "%PDF" in rej[0])

_, rej = validate_uploads([Up("fake.docx", b"%PDF-1.7 not a docx")])
check("PDF bytes with .docx name rejected", len(rej) == 1 and "real DOCX" in rej[0])

_, rej = validate_uploads([Up("payload.exe", b"MZ\x90\x00")])
check("disallowed extension rejected", len(rej) == 1 and "PDF and DOCX" in rej[0])

ups, rej = validate_uploads([Up("../../etc/passwd.pdf", real_pdf)])
check("traversal name sanitised", len(rej) == 0 and ups[0].name == "passwd.pdf",
      ups[0].name if ups else "-")

big = b"%PDF-1.7" + b"0" * int((MAX_FILE_MB + 1) * 1024 * 1024)
_, rej = validate_uploads([Up("huge.pdf", big)])
check("oversize (>%dMB) rejected" % MAX_FILE_MB, len(rej) == 1 and "exceeds" in rej[0])
del big

many = [Up("c%d.pdf" % i, real_pdf) for i in range(_app.MAX_FILES + 5)]
ups, rej = validate_uploads(many)
check("file-count budget enforced",
      len(ups) == _app.MAX_FILES and any("Batch limit" in r for r in rej),
      "%d files" % len(ups))

dup = validate_uploads([Up("same.pdf", real_pdf), Up("same.pdf", real_pdf)])[0]
check("duplicate names de-duplicated",
      len(dup) == 2 and dup[0].name != dup[1].name,
      "%s / %s" % (dup[0].name, dup[1].name))

print("\n=== 3. text extraction: PDF, DOCX, encrypted, corrupt ===")
t = read_upload(ups[0])
check("real PDF extracts text", len(t) > 80, "%d chars" % len(t))

try:
    read_pdf_bytes(b"%PDF-1.7\ngarbage%%EOF")
    check("corrupt PDF raises CVReadError", False)
except CVReadError as e:
    check("corrupt PDF raises CVReadError", True, str(e)[:70])

from pypdf import PdfWriter  # noqa: E402

_w = PdfWriter()
_w.add_blank_page(width=200, height=200)
_w.encrypt("secret123")
_buf = io.BytesIO()
_w.write(_buf)
try:
    read_pdf_bytes(_buf.getvalue())
    check("password-protected PDF raises CVReadError", False)
except CVReadError as e:
    check("password-protected PDF raises CVReadError", "password" in str(e), str(e)[:70])

import docx as _docx  # noqa: E402

_d = _docx.Document()
_d.add_paragraph("Site Management Field Assistant with five years experience.")
_tbl = _d.add_table(rows=1, cols=2)
_tbl.rows[0].cells[0].text = "Skill"
_tbl.rows[0].cells[1].text = "Years"
_dt = io.BytesIO()
_d.save(_dt)
check("DOCX extracts text + tables", "Site Management" in read_docx_bytes(_dt.getvalue()))

_bad = io.BytesIO()
_bad.write(b"PK\x03\x04" + b"\x00" * 64)
try:
    read_docx_bytes(_bad.getvalue())
    check("corrupt DOCX raises CVReadError", False)
except CVReadError:
    check("corrupt DOCX raises CVReadError", True)

print("\n=== 4. filters, bonus, education gate, remarks ===")
check("exclude-keyword filter fires", keyword_filters("posted in dhaka", ["dhaka"], []) is not None)
check("require-keyword filter fires", keyword_filters("general worker", [], ["site management"]) is not None)
check("clean CV passes filters",
      keyword_filters("site management officer in coxsbazar", ["dhaka"], ["site management"]) is None)
check("bonus adds per criterion and is capped",
      bonus_points("site management plus report writing",
                   ["site management", "report writing"], 2.0, 3.0)[0] == 3.0
      and bonus_points("site management",
                       ["site management", "report writing"], 2.0, 3.0)[0] == 2.0)
check("split_keywords", split_keywords("a, b\nc; d") == ["a", "b", "c", "d"])
check("remark is structured", "STRONG FIT" in recruiter_remark(78.0, ["Strong match."], ["No gap."]))
check("exclusion remark explains", "EXCLUDED" in exclusion_remark("keyword"))
check("error remark explains", "NOT EVALUATED" in error_remark("encrypted"))
_v1 = education_verdict("Bachelor of Social Science BSS degree", ["bachelor"])
check("education gate (met)", bool(_v1) and _v1[0] is True, str(_v1))
_v2 = education_verdict("Secondary school diploma SSC", ["masters"])
check("education gate (below)", bool(_v2) and _v2[0] is False, str(_v2))

print("\n=== 5. end-to-end screening (the app's own loop) ===")
profile = build_profile(jd_text())
check("JD profile parsed from text", profile is not None and profile.title != "",
      getattr(profile, "title", "-"))

work = copy.deepcopy(profile)
work.min_experience_years, work.max_experience_years = 3.0, 5.0
work.education_levels = ["bachelor"]

entries, real_blacklist = blacklist_entries()
sample = cv_fixtures()
rows = []
for up in sample:
    try:
        text = read_upload(up)
    except CVReadError as e:
        rows.append(_app._row(up.name, 0.0, "READ ERROR", error_remark(str(e)), str(e)))
        continue
    verdict = ev.blacklist_check(text, up.name, entries, "flag")
    if verdict.excluded:
        rows.append(_app._row(up.name, 0.0, "EXCLUDED",
                              exclusion_remark(verdict.excluded), verdict.excluded))
        continue
    res = ev.score_candidate(text, work)
    bonus, hits = bonus_points(text.lower(), ["site management", "community mobilization"], 1.0, 6.0)
    final = round(min(100.0, res.score + bonus), 1)
    final_res = dataclasses.replace(res, score=final)
    rows.append(_app._row(up.name, final,
                          "REVIEW" if verdict.review else "RANKED",
                          recruiter_remark(final, res.strengths, res.gaps, hits),
                          ev.build_reasoning(final_res), res.strengths, res.gaps, bonus))

check("every sample CV scored or explained", len(rows) == len(sample),
      "%d/%d" % (len(rows), len(sample)))
check("scores within 0-100", all(0 <= r["Match_Score"] <= 100 for r in rows))
check("scorecards carry strengths",
      all(r["Strengths"] or r["Status"] != "RANKED" for r in rows))
check("scorecards carry a remark",
      all("score" in r["Recruiter_Remark"].lower() for r in rows))

if not real_blacklist:
    _strong = next((r for r in rows if r["Candidate_File_Name"].startswith("md-rahim")), None)
    _weak = next((r for r in rows if r["Candidate_File_Name"].startswith("karim")), None)
    check("relevant CV outranks an unrelated one",
          bool(_strong and _weak) and _strong["Match_Score"] > _weak["Match_Score"],
          "strong=%s weak=%s" % (_strong and _strong["Match_Score"],
                                 _weak and _weak["Match_Score"]))

frame = export_frame(rows)
check("export frame columns",
      list(frame.columns) == ["Rank", "Candidate_File_Name", "Match_Score",
                              "Bonus_Points", "Status", "Strengths", "Weaknesses",
                              "Recruiter_Remark", "Summary_Reasoning"],
      ",".join(frame.columns))
csv_b = frame.to_csv(index=False).encode("utf-8-sig")
check("CSV export bytes", len(csv_b) > 200 and csv_b[:3] == b"\xef\xbb\xbf", "%d B" % len(csv_b))
_xb = io.BytesIO()
frame.to_excel(_xb, index=False, engine="openpyxl", sheet_name="Candidates")
check("Excel export bytes", len(_xb.getvalue()) > 1000, "%d B" % len(_xb.getvalue()))

print("\n=== 6. blacklist gate (phone => exclude, name => review only) ===")
e0 = entries[0]
check("blacklist has at least one entry", bool(entries), "" if real_blacklist else "(synthetic)")
_v = ev.blacklist_check("Contact me on %s anytime. " % e0.phones[0] + "x" * 300, "c.pdf", entries, "flag")
check("phone match excludes", _v.excluded is not None, str(_v.excluded)[:60])
_nm = e0.names[0] if e0.names else "zz yy"
_v2 = ev.blacklist_check("References: %s was my supervisor. " % _nm + "x" * 400,
                         "other.pdf", entries, "flag")
check("name-only match only flags",
      _v2.excluded is None and _v2.review is not None, str(_v2.review)[:70])

print("\n=== 7. the Streamlit script itself boots ===")
_at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=240)
_at.run()
check("no exception on boot", not _at.exception,
      str([str(e.value)[:200] for e in _at.exception]))
check("no st.error on boot", not _at.error, str([m.value[:120] for m in _at.error]))

print("\n" + "=" * 70)
print("RESULT:", "ALL PASS" if not fails else "%d FAILED -> %s" % (len(fails), fails))
sys.exit(1 if fails else 0)
